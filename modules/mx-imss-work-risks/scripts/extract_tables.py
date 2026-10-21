"""Extract tables from the IMSS statistical yearbooks (Memoria Estadistica), chapter VII.

The Mexican Social Security Institute (IMSS) publishes a statistical yearbook each year. Its
chapter VII, "Salud en el Trabajo" (occupational health), is an Excel workbook of about 50 tables
on work risks among workers covered by IMSS's work risk insurance. This script reads the 2016 to
2024 workbooks (downloaded by scripts/fetch_data.R) and writes, to data/tidy/:
  national.csv               companies, workers, cases, permanent incapacities and deaths by type
                             of risk, 2011 to 2024
  by_state.csv               the same by IMSS state office (OOAD, formerly delegacion), 2016 to 2024
  by_activity.csv            workers and cases by division of economic activity and type of risk
  by_age_sex.csv             workers and cases by age group and sex
  occupational_diseases.csv  occupational disease cases by diagnosis, 2011 to 2024
  accident_causes.csv        work and commuting accidents by external cause, sex and age group
  body_region_injury.csv     work accidents by body region, type of injury and sex
  occupations.csv            work accidents in the most affected occupations, by type of injury
                             and sex
Spanish labels are kept, and English ones added from scripts/labels_es_en.csv.

Tables are found by their title, since table numbers change between yearbooks. Merged header
cells are filled in, so each value column gets a path of headers (for example "Accidentes de
Trabajo" / "Casos"); published rates are dropped, since their units change between yearbooks, and
the notebooks compute rates from the counts and workers.

Needs openpyxl (pip install openpyxl). Written in Python, like the extraction scripts of other
modules; the outputs are committed, so you only need to run this to update them.

Usage (from the module folder):  python scripts/extract_tables.py
"""
import csv
import re
import unicodedata
from pathlib import Path

import openpyxl

module = Path(__file__).resolve().parent.parent
raw = module / "data" / "raw"
tidy = module / "data" / "tidy"
tidy.mkdir(exist_ok=True)
YEARS = range(2016, 2025)


# ---- Reading a table ---------------------------------------------------------------------------
def clean(s):
    if not isinstance(s, str):
        return s
    return re.sub(r"\s+", " ", s.replace("\n", " ")).strip()


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def key(s):
    """Label key for matching: no accents, footnote marks, final periods or case."""
    s = unicodedata.normalize("NFKD", clean(s)).encode("ascii", "ignore").decode()
    s = re.sub(r"\(\d\)|\*", "", s)
    s = re.sub(r"\s*\(", " (", s)
    s = re.sub(r"\s+", " ", s).strip(" .").lower()
    # Some headers are spaced out letter by letter ("T o t a l N a c i o n a l")
    if re.fullmatch(r"(\w ){3,}\w", s):
        s = {"totalnacional": "total nacional"}.get(s.replace(" ", ""), s.replace(" ", ""))
    return s


def find(wb, pattern):
    """The sheet whose title (the line after "Cuadro No.") matches pattern."""
    for ws in wb.worksheets:
        if not ws.title.strip().upper().startswith("VII"):
            continue
        texts = [clean(c) for r in ws.iter_rows(values_only=True, max_row=8) for c in r
                 if isinstance(c, str) and c.strip()]
        texts = [t for t in texts if not t.startswith(("Regresar", "Cuadro"))]
        if texts and re.search(pattern, texts[0], re.I):
            return ws
    return None


def read_table(ws):
    """Return {column: header path} and a list of (row label, {column: value})."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    width = max(len(r) for r in grid)
    grid = [r + [None] * (width - len(r)) for r in grid]
    for m in ws.merged_cells.ranges:
        v = grid[m.min_row - 1][m.min_col - 1]
        for i in range(m.min_row - 1, min(m.max_row, len(grid))):
            for j in range(m.min_col - 1, min(m.max_col, width)):
                grid[i][j] = v
    grid = [[clean(c) for c in r] for r in grid]
    title_row = next(i for i, r in enumerate(grid) if any(
        isinstance(c, str) and len(c) > 30 and not c.startswith(("Regresar", "Cuadro"))
        for c in r))
    start = next(i for i, r in enumerate(grid) if i > title_row and sum(map(is_num, r)) >= 2)
    numeric = sorted({j for r in grid[start:] for j, c in enumerate(r) if is_num(c)})
    # Some workbooks lost their merged cells: fill a blank header cell from its left neighbour
    # when both sit under the same header in the rows above
    header = [list(r) for r in grid[title_row + 1:start]]
    for i, r in enumerate(header):
        for j in range(1, width):
            above = [header[k][j] for k in range(i)] == [header[k][j - 1] for k in range(i)]
            if r[j] in (None, "") and r[j - 1] not in (None, "") and above and j > min(numeric):
                r[j] = r[j - 1]
    paths = {}
    for j in numeric:
        parts = []
        for r in header:
            v = r[j]
            if v not in (None, "") and str(v) not in parts:
                parts.append(str(v))
        paths[j] = tuple(parts)
    rows = []
    for r in grid[start:]:
        labels = [c for c in r[:min(numeric)] if isinstance(c, str) and c]
        label = labels[-1] if labels else ""
        if re.match(r"^(\(\d+\)|\d\)|Fuente|Nota|\*)", label):
            break
        vals = {j: r[j] for j in numeric if is_num(r[j])}
        if label or vals:
            rows.append((label, vals))
    return paths, rows


def tables(pattern, years=YEARS):
    """Yield (year, paths, rows) for the table matching pattern in each yearbook."""
    for year in years:
        wb = openpyxl.load_workbook(raw / f"memoria_{year}_salud_trabajo.xlsx", data_only=True)
        ws = find(wb, pattern)
        if ws is None:
            raise ValueError(f"{year}: no table matches {pattern!r}")
        yield (year, *read_table(ws))


def year_columns(rows):
    """For tables with years as columns: {column: year} from the first row of years."""
    for label, vals in rows:
        if vals and all(isinstance(v, int) and 2000 < v < 2100 for v in vals.values()):
            seen, out = set(), {}
            for j, v in sorted(vals.items()):
                if v not in seen:  # the first column of each year (cases, not percent)
                    out[j] = v
                    seen.add(v)
            return out
    raise ValueError("no row of years")


RISKS = [("sin covid", "Occupational diseases excluding COVID-19"),
         ("respiratorias asociadas a covid", "COVID-19-related respiratory diseases"),
         ("covid", "COVID-19"),
         ("accidentes y enfermedades", "Work accidents and diseases"),
         ("trayecto", "Commuting accidents"), ("enfermedad", "Occupational diseases"),
         ("accidentes de trabajo", "Work accidents"), ("riesgo", "All work risks")]


def risk_of(text):
    t = key(text)
    r = next((v for k, v in RISKS if k in t), None)
    if r is None:
        raise ValueError("unknown risk: " + text)
    return r


def national_name(label):
    return "National" if key(label) in ("total", "nacional", "total nacional") else label


# ---- National series ---------------------------------------------------------------------------
national = {}
for year, paths, rows in tables(r"^Empresas, trabajadores, riesgos de trabajo, incapacidades"):
    cols = year_columns(rows)
    section = None
    for label, vals in rows:
        k = key(label)
        if "por cada" in k:
            break  # rates follow
        if k.startswith("empresas"):
            measure, risk = "Companies", ""
        elif k.startswith("trabajadores"):
            measure, risk = "Workers", ""
        elif k in ("riesgo de trabajo", "riesgos de trabajo", "incapacidades permanentes",
                   "defunciones"):
            section = {"r": "Cases", "i": "Permanent incapacities", "d": "Deaths"}[k[0]]
            measure, risk = section, "All work risks"
        elif section and vals:
            measure, risk = section, risk_of(label)
        else:
            continue
        for j, y in cols.items():
            if j in vals:
                # each year from the latest yearbook that has it
                national[(y, measure, risk)] = {"year": y, "measure": measure, "risk": risk,
                                                "value": vals[j], "source_yearbook": year}

# ---- By state ----------------------------------------------------------------------------------
by_state = {}
for pattern, measure in [(r"^Empresas, trabajadores, riesgos de trabajo e indicadores", "cases"),
                         (r"^Incapacidades permanentes ?\(1\), trabajadores", "permanent_incapacities"),
                         (r"^Defunciones por riesgo de trabajo e indicadores", "deaths")]:
    for year, paths, rows in tables(pattern):
        for label, vals in rows:
            if not vals:
                continue
            state = national_name(re.sub(r"\s*(\(\d\))+$", "", label))
            for j, v in vals.items():
                p = paths[j]
                last = key(p[-1]) if p else ""
                if key(p[0]).startswith("numero de empresas"):
                    by_state.setdefault((year, state, ""), {})["companies"] = v
                elif key(p[0]).startswith("trabajadores"):
                    by_state.setdefault((year, state, ""), {})["workers"] = v
                elif "casos" in last or last in ("incapacidades", "defunciones"):
                    risk = risk_of(" ".join(p))
                    by_state.setdefault((year, state, risk), {})[measure] = v

# ---- By activity, age and sex --------------------------------------------------------------------
by_activity = []
for year, paths, rows in tables(r"divisi.n de actividad econ.mica ?\(1\) y tipo de riesgo"):
    for label, vals in rows:
        if not vals:
            continue
        rec = {"year": year, "activity_es": national_name(label)}
        for j, v in vals.items():
            p = paths[j]
            if key(p[0]).startswith("trabajadores"):
                rec["workers"] = v
            elif key(p[-1]) == "casos":
                rec[risk_of(p[0])] = v
        by_activity.append(rec)

by_age_sex = []
for year, paths, rows in tables(r"^Trabajadores bajo seguro de riesgo de trabajo, riesgo de trabajo y tasa"):
    for label, vals in rows:
        for j, v in vals.items():
            group, sex = paths[j][0], paths[j][-1]
            if "por cada" in key(group) or key(sex) == "total":
                continue
            measure = "workers" if key(group).startswith("trabajadores") else "cases"
            by_age_sex.append({"year": year, "age_group": re.sub(r"\s*-\s*", "-", label),
                               "sex": {"hombres": "Male", "mujeres": "Female"}[key(sex)],
                               "measure": measure, "value": v})

# ---- Occupational diseases -----------------------------------------------------------------------
diseases = {}
for year, paths, rows in tables(r"^Enfermedades de trabajo, seg.n naturaleza de la lesi.n ?\(1\), ?20"):
    cols = year_columns(rows)
    for label, vals in rows:
        if not vals or all(isinstance(v, int) and 2000 < v < 2100 for v in vals.values()):
            continue
        for j, y in cols.items():
            if j in vals:
                diseases[(y, key(label))] = {"year": y, "disease_es": national_name(label),
                                             "cases": vals[j], "source_yearbook": year}

# ---- Accident detail -----------------------------------------------------------------------------
SEX = {"h": "Male", "m": "Female"}

causes = []
for kind, pattern in [("Work accidents", r"^Accidentes de trabajo.{0,8}seg.n causa externa.{0,6}, sexo y grupos"),
                      ("Commuting accidents", r"^Accidentes en trayecto.{0,8}seg.n causa externa.{0,6}, sexo y grupos")]:
    years = YEARS if kind == "Work accidents" else range(2019, 2025)
    for year, paths, rows in tables(pattern, years):
        for label, vals in rows:
            if key(label) == "total":
                continue
            for j, v in vals.items():
                p = paths[j]
                if len(p) < 2 or key(p[-1]) not in SEX or key(p[-2]).startswith("total"):
                    continue
                causes.append({"year": year, "risk": kind, "cause_es": label,
                               "sex": SEX[key(p[-1])], "age_group": p[-2], "cases": v})


def injury_by_sex(pattern, row_name):
    out = []
    for year, paths, rows in tables(pattern):
        for label, vals in rows:
            if key(label) in ("total", "total nacional", "%"):
                continue
            for j, v in vals.items():
                p = paths[j]
                if len(p) < 2 or key(p[-1]) not in SEX or key(p[-2]).startswith("total"):
                    continue
                out.append({"year": year, row_name: label, "injury_es": p[-2],
                            "sex": SEX[key(p[-1])], "cases": v})
    return out


body = injury_by_sex(r"^Accidentes de trabajo.{0,8}seg.n regi.n anat.mica, tipo de lesi", "body_region_es")
occupations = injury_by_sex(r"^Accidentes de trabajo.{0,8}seg.n ocupaci.n, tipo de lesi", "occupation_es")

# ---- Labels --------------------------------------------------------------------------------------
labels, canonical = {}, {}
with open(Path(__file__).resolve().parent / "labels_es_en.csv", encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        labels[(r["variable"], key(r["spanish"]))] = r["english"]
        canonical.setdefault((r["variable"], r["english"]), r["spanish"])
missing = set()


def translate(variable, text):
    """English and Spanish labels; the Spanish one is spelled the same way in every year."""
    text = re.sub(r"(\(\d\))+$", "", clean(text)).strip(" .")
    if text == "National":
        return "National", "Nacional"
    english = labels.get((variable, key(text)))
    if english is None:
        if variable in ("state", "age_group"):  # names that need no translation
            return text, text
        missing.add((variable, text))
        return "", text
    return english, canonical[(variable, english)]


# The 35 IMSS state offices: Mexico City, the State of Mexico and Veracruz have two each
STATE_OF = {"Mexico City North": "Mexico City", "Mexico City South": "Mexico City",
            "State of Mexico East": "State of Mexico", "State of Mexico West": "State of Mexico",
            "Veracruz North": "Veracruz", "Veracruz South": "Veracruz"}

published = {(y, measure, r) for (y, _, r), m in by_state.items() for measure in m}
out_state = []
for (year, state, risk), m in sorted(by_state.items(), key=lambda x: (x[0][0], x[0][1] != "National", x[0][1], x[0][2])):
    if risk == "":
        continue
    office, office_es = translate("state", state)
    base = by_state.get((year, state, ""), {})
    # A blank cell is zero when the table has that column, and missing when it does not
    def value(measure):
        return m.get(measure, 0 if (year, measure, risk) in published else None)
    out_state.append({"year": year, "office": office, "office_es": office_es,
                      "state": STATE_OF.get(office, office), "companies": base.get("companies"),
                      "workers": base.get("workers"), "risk": risk, "cases": value("cases"),
                      "permanent_incapacities": value("permanent_incapacities"),
                      "deaths": value("deaths")})

out_activity = []
for r in by_activity:
    activity, activity_es = translate("activity", r["activity_es"])
    for risk in [v for _, v in RISKS]:
        if risk in r:
            out_activity.append({"year": r["year"], "activity": activity,
                                 "activity_es": activity_es, "workers": r.get("workers"),
                                 "risk": risk, "cases": r[risk]})

wide = {}
for r in by_age_sex:
    if key(r["age_group"]) == "total":
        continue
    age = translate("age_group", r["age_group"])[0]
    wide.setdefault((r["year"], age, r["sex"]), {})[r["measure"]] = r["value"]
out_age = [{"year": y, "age_group": a, "sex": s, "workers": m.get("workers", 0),
            "cases": m.get("cases", 0)} for (y, a, s), m in wide.items()]

merged = {}
for r in diseases.values():
    if r["disease_es"] == "National":
        disease, disease_es = "Total", "Total"
    else:
        disease, disease_es = translate("disease", r["disease_es"])
    # The 2024 yearbook lists De Quervain's tenosynovitis twice (one row is all zeros)
    m = merged.setdefault((r["year"], disease), {"year": r["year"], "disease": disease,
                                                 "disease_es": disease_es, "cases": 0,
                                                 "source_yearbook": r["source_yearbook"]})
    m["cases"] += r["cases"]
out_diseases = sorted(merged.values(), key=lambda r: (r["year"], r["disease"] != "Total", -r["cases"]))

out_causes = []
for r in causes:
    cause, cause_es = translate("cause", r["cause_es"])
    out_causes.append({"year": r["year"], "risk": r["risk"], "cause": cause, "cause_es": cause_es,
                       "sex": r["sex"], "age_group": translate("age_group", r["age_group"])[0],
                       "cases": r["cases"]})


def injuries(records, row_var):
    out = []
    for r in records:
        row, row_es = translate(row_var, r[row_var + "_es"])
        injury, injury_es = translate("injury", r["injury_es"])
        out.append({"year": r["year"], row_var: row, row_var + "_es": row_es, "injury": injury,
                    "injury_es": injury_es, "sex": r["sex"], "cases": r["cases"]})
    return out


out_body = injuries(body, "body_region")
out_occupations = injuries(occupations, "occupation")

if missing:
    raise ValueError("No English label in scripts/labels_es_en.csv for: " +
                     "; ".join(f"{v}: {t}" for v, t in sorted(missing)))

# ---- Checks --------------------------------------------------------------------------------------
# Parts should add up to totals. IMSS's own tables differ by a few cases in some years; those
# are listed but kept as published. A difference of more than 0.5% stops the script.
problems, small = [], []


def check(what, a, b, tolerance=0.005):
    if a == b:
        return
    (problems if abs(a - b) > tolerance * b else small).append(f"  {what}: {a} vs {b}")


nat = {(r["year"], r["risk"]): r for r in out_state if r["office"] == "National"}
for (year, risk), r in nat.items():
    offices = [x for x in out_state
               if x["year"] == year and x["risk"] == risk and x["office"] != "National"]
    for m in ("cases", "permanent_incapacities", "deaths"):
        if r[m] is not None:
            check(f"{year} {risk} {m}: offices vs national",
                  sum(x[m] or 0 for x in offices), r[m])
    acts = [x for x in out_activity
            if x["year"] == year and x["risk"] == risk and x["activity"] != "National"]
    if acts:
        check(f"{year} {risk}: activities vs national", sum(x["cases"] for x in acts), r["cases"])
for year in YEARS:
    check(f"{year} workers: offices vs national",
          sum(x["workers"] for x in out_state if x["year"] == year and x["risk"] == "All work risks"
              and x["office"] != "National"), nat[(year, "All work risks")]["workers"])
    check(f"{year} age and sex: cases vs national",
          sum(x["cases"] for x in out_age if x["year"] == year), nat[(year, "All work risks")]["cases"])
    for kind in ("Work accidents", "Commuting accidents"):
        c = [x["cases"] for x in out_causes if x["year"] == year and x["risk"] == kind]
        if c:
            check(f"{year} {kind}: causes vs national", sum(c), nat[(year, kind)]["cases"])
    for name, rows in (("body regions", out_body), ("occupations", out_occupations)):
        check(f"{year} work accidents: {name} vs national",
              sum(x["cases"] for x in rows if x["year"] == year), nat[(year, "Work accidents")]["cases"])
for year in sorted({r["year"] for r in out_diseases}):
    rows = [r for r in out_diseases if r["year"] == year]
    parts = sum(r["cases"] for r in rows if r["disease"] != "Total")
    total = sum(r["cases"] for r in rows if r["disease"] == "Total")
    # The total includes cases with no diagnosis recorded, so the parts can fall short of it
    check(f"{year} diseases: parts vs total", max(parts, total), total)
# The national series takes each year from the latest yearbook; the state table from that year's
# own yearbook. Differences are revisions.
revisions = []
for (year, measure, risk), r in sorted(national.items()):
    if year in YEARS and risk and measure != "Companies":
        own = nat.get((year, risk), {}).get({"Cases": "cases", "Deaths": "deaths",
                                             "Permanent incapacities": "permanent_incapacities"}[measure])
        if own is not None and own != r["value"]:
            revisions.append(f"  {year} {measure}, {risk}: {own} in the {year} yearbook, "
                             f"{r['value']} in the {r['source_yearbook']} yearbook")


# ---- Write ----------------------------------------------------------------------------------------
def write(name, rows):
    with open(tidy / name, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


write("national.csv", sorted(national.values(), key=lambda r: (r["measure"], r["risk"], r["year"])))
write("by_state.csv", out_state)
write("by_activity.csv", out_activity)
write("by_age_sex.csv", out_age)
write("occupational_diseases.csv", out_diseases)
write("accident_causes.csv", out_causes)
write("body_region_injury.csv", out_body)
write("occupations.csv", out_occupations)
for f in sorted(tidy.glob("*.csv")):
    print(f"{f.name:28} {sum(1 for _ in open(f, encoding='utf-8')) - 1:7} rows")
print("Revised figures (state table in its own yearbook vs the latest yearbook):")
print("\n".join(revisions) or "  none")
print("Small differences in the published tables (kept as published):")
print("\n".join(small) or "  none")
if problems:
    raise ValueError("Checks failed:\n" + "\n".join(problems))
print("All checks passed")
