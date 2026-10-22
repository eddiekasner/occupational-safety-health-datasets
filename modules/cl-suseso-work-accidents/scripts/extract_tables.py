"""Extract work accident statistics from SUSESO's annual "Estadisticas de la Seguridad Social".

Chile's Superintendencia de Seguridad Social (SUSESO) publishes a statistical yearbook each year as
an Excel workbook. Its chapter I covers the insurance against work accidents and occupational
diseases of Law 16.744: protected workers, work accidents, commuting accidents, occupational
diseases, accident rates, days lost and deaths, by insurer (the three private mutual insurers and
the public Instituto de Seguridad Laboral), economic sector, region, sex and size of employer. This
script reads the 2014 to 2025 yearbooks (downloaded by scripts/fetch_data.R) and writes, to
data/tidy/:
  national.csv     national series by insurer (each table covers the last 2 to 5 years; each year
                   is taken from the latest yearbook that has it)
  by_sector.csv    by economic sector and insurer
  by_region.csv    by region and insurer
  by_sex.csv       by sex, and by sex and sector
  by_size.csv      by size of employer and insurer
Tables are found by their titles, because their numbers change between yearbooks.

Needs openpyxl (pip install openpyxl). Written in Python, like the extraction scripts of other
modules; the outputs are committed, so you only need to run this to update them.

Usage (from the module folder):  python scripts/extract_tables.py
"""
import csv
import re
import unicodedata
from pathlib import Path

import openpyxl

module = Path(__file__).resolve().parent.parent
raw = module / "data" / "raw"
tidy = module / "data" / "tidy"
tidy.mkdir(exist_ok=True)
YEARS = range(2014, 2026)


def clean(s):
    return re.sub(r"\s+", " ", s).strip() if isinstance(s, str) else s


def key(s):
    """Text for matching: no accents, footnote marks or case."""
    s = unicodedata.normalize("NFKD", clean(str(s))).encode("ascii", "ignore").decode()
    s = re.sub(r"\(\d\)|\(\*+\)|\*", "", s)
    s = re.sub(r"\s+", " ", s).strip(" .:").lower()
    # Some labels are spaced out letter by letter ("T O T A L")
    return s.replace(" ", "") if re.fullmatch(r"(\w ){2,}\w", s) else s


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def sheet_grid(ws):
    """The cells of a sheet, with merged cells filled."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    width = max(len(r) for r in grid)
    grid = [r + [None] * (width - len(r)) for r in grid]
    for m in ws.merged_cells.ranges:
        v = grid[m.min_row - 1][m.min_col - 1]
        for i in range(m.min_row - 1, min(m.max_row, len(grid))):
            for j in range(m.min_col - 1, min(m.max_col, width)):
                grid[i][j] = v
    return grid


def more_tables(grid):
    """Tables below the first one: (title, grid of the table). Some are side by side."""
    out = []
    for i, r in enumerate(grid):
        # (the first cell of each, since merged cells repeat)
        cols = [j for j, c in enumerate(r) if isinstance(c, str) and re.match(r"(?i)^\s*cuadro", c)
                and (j == 0 or r[j - 1] != c)]
        if i < 3 or not cols:
            continue
        for a, b in zip(cols, cols[1:] + [len(r)]):
            sub = [row[a:b] for row in grid[i:]]
            title = next((clean(c) for row in sub[:4] for c in row if isinstance(c, str)
                          and len(c.strip()) > 30 and not re.match(r"(?i)cuadro|volver", c.strip())), "")
            out.append((title, sub))
    return out


def read_table(grid):
    """Return the title, {column: header path} and rows of (section, label, {column: value})."""
    width = len(grid[0])
    texts = [(i, clean(c)) for i, r in enumerate(grid[:10]) for c in r if isinstance(c, str) and c.strip()]
    title_row, title = next(((i, t) for i, t in texts if len(t) > 30
                             and not re.match(r"(?i)cuadro|volver", t)), (0, ""))
    # The header: the rows between the title and the first row with two or more numbers that are
    # not years
    def data_like(r):
        return sum(is_num(c) and not (isinstance(c, int) and 1990 < c < 2100) for c in r) >= 2
    start = next(i for i, r in enumerate(grid) if i > title_row and data_like(r))
    # The table ends at its notes (some sheets have a second table below)
    end = next((i for i in range(start, len(grid)) if any(
        isinstance(c, str) and re.match(r"(?i)^\s*(fuente|nota|cuadro)", c) for c in grid[i])), len(grid))
    grid = grid[:end]
    numeric = sorted({j for r in grid[start:] for j, c in enumerate(r)
                      if is_num(c) and not (isinstance(c, int) and 1990 < c < 2100)})
    label_cols = [j for j in range(min(numeric))]
    header = [list(r) for r in grid[title_row + 1:start]]
    paths = {}
    for j in numeric:
        parts = []
        for r in header:
            v = r[j]
            if v not in (None, "") and str(v) not in parts and not re.match(r"(?i)volver", str(v)):
                parts.append(clean(str(v)))
        paths[j] = tuple(parts)
    rows, section = [], ""
    # The row just above the first data row may name the first section
    first = start - 1 if start and not any(is_num(c) for c in grid[start - 1]) else start
    for r in grid[first:]:
        labels = [clean(r[j]) for j in label_cols if isinstance(r[j], str) and r[j].strip()]
        values = {j: r[j] for j in numeric if is_num(r[j])}
        label = labels[-1] if labels else ""
        if re.match(r"(?i)^(fuente|nota|\(\d|\*|\d\))", label):
            break
        if not values:
            # Headings of sections are in capitals; other rows without numbers are empty rows
            letters = [c for c in label if c.isalpha()]
            if letters and sum(c.isupper() for c in letters) / len(letters) > 0.8:
                section = label
            continue
        rows.append((section, label or "TOTAL", values))  # (one total row has no label)
    return title, paths, rows


# ---- The tables, found by their titles ---------------------------------------------------------
# (output table, dimension of the rows, title pattern). Titles are matched without accents or case.
INSURER = r"(organismos? administrador(es)?|mutual|mutuales|mutualidad|mutualidades)"
TABLES = [
    ("national", "insurer", r"^numero promedio (mensual )?de (trabajadores|personas unicas) protegid\w*( \(\d\))? por el seguro de la ley n. ?16\.?744( \(\d\))?$"),
    ("national", "insurer", r"^numero de accidentes .*(segun|por) " + INSURER + r"$"),
    ("national", "insurer", r"^tasas de accidentabilidad .*segun " + INSURER + r"$"),
    ("national", "insurer", r"^numero de dias perdidos .*segun " + INSURER + r"$"),
    ("national", "insurer", r"^numero promedio de dias perdidos .*segun " + INSURER + r"$"),
    ("national", "insurer", r"^numero de fallecidos .*segun (fecha del fallecimiento, )?tipo de accidente y organismo administrador$"),
    ("national", "insurer", r"^tasa de mortalidad segun tipo de accidente y organismo administrador$"),
    ("sector", "sector", r"^numero promedio mensual de (trabajadores|relaciones laborales) protegid.*actividad economica y " + INSURER + r"$"),
    ("sector", "sector", r"^numero de accidentes .*actividad economica y " + INSURER + r"$"),
    ("sector", "sector", r"^tasas de accidentabilidad .*actividad economica y " + INSURER + r"$"),
    ("sector", "sector", r"^numero de dias perdidos .*actividad economica y " + INSURER + r"$"),
    ("sector", "sector", r"^numero promedio de dias perdidos .*actividad economica y " + INSURER + r"$"),
    ("sector", "sector", r"^numero de fallecidos .*tipo de accidente,? (por )?actividad economica y organismo administrador$"),
    ("sector", "sector", r"^tasa de mortalidad .*actividad economica"),
    ("region", "region", r"^numero promedio mensual de (trabajadores|relaciones laborales) protegid.*region ?,? (y |segun )?" + INSURER + r"$"),
    ("region", "region", r"^numero de dias perdidos .*region y " + INSURER + r"$"),
    ("region", "region", r"^numero promedio de dias perdidos .*region y " + INSURER + r"$"),
    ("region", "region", r"^numero de accidentes .*region y " + INSURER + r"$"),
    ("region", "region", r"^numero de fallecidos .*tipo de accidente,? region y organismo administrador$"),
    ("sex", "insurer", r"^numero promedio mensual de (trabajadores|relaciones laborales) protegid.*segun (sexo y organismo administrador|organismo administrador y sexo)$"),
    ("sex", "sector", r"^numero promedio mensual de (trabajadores|relaciones laborales) protegid.*segun (sexo y actividad economica|organismo administrador, actividad economica y sexo)$"),
    ("sex", "insurer", r"^numero de dias perdidos .*segun " + INSURER + r" y sexo$"),
    ("sex", "insurer", r"^numero de fallecidos .*tipo de accidente, organismo administrador y sexo$"),
    ("sex", "insurer", r"^numero de accidentes.*segun " + INSURER + r" y sexo$"),
    ("sex", "sector", r"^numero de accidentes .*actividad economica y sexo$"),
    ("sex", "sector", r"^numero de fallecidos .*actividad economica y sexo$"),
    ("sex", "region", r"^numero de fallecidos .*region y sexo$"),
    ("sex", "region", r"^numero de accidentes .*region y sexo$"),
    ("size", "size", r"^numero promedio mensual de (trabajadores|relaciones laborales) protegid.*tamano de la entidad empleadora"),
    ("size", "size", r"^numero de accidentes .*tamano de la entidad empleadora"),
    ("size", "size", r"^numero de dias perdidos .*tamano de la entidad empleadora"),
    ("size", "size", r"^numero promedio de dias perdidos .*tamano de la entidad empleadora"),
]

# Occupational diseases by date of notification (the criterion the 2025 yearbook uses elsewhere)
OD_TABLE = r"^numero de enfermedades profesionales segun fecha de denuncia por " + INSURER


def chapter_one(wb):
    """The sheets of chapter I (work accidents and occupational diseases)."""
    out = []
    for ws in wb.worksheets:
        if re.match(r"(?i)cap.tulo ii\b", ws.title.strip()):
            break
        out.append(ws)
    return out


long = []
for year in YEARS:
    wb = openpyxl.load_workbook(raw / f"estadisticas_seguridad_social_{year}.xlsx", data_only=True)
    for ws in chapter_one(wb):
        head = [clean(c) for r in ws.iter_rows(values_only=True, max_row=8) for c in r
                if isinstance(c, str) and len(c.strip()) > 30 and not re.match(r"(?i)cuadro|volver", c.strip())]
        if not head:
            continue
        match = [(out, dim) for out, dim, pat in TABLES if re.search(pat, key(head[0]))]
        if not match:
            continue
        out, dim = match[0]
        grid = sheet_grid(ws)
        tables = [(out, dim, grid)]
        # The 2025 yearbook puts occupational diseases in a second table on the sheet of accidents
        tables += [(out, dim, sub) for t, sub in more_tables(grid) if re.search(OD_TABLE, key(t))]
        for out, dim, g in tables:
            title, paths, rows = read_table(g)
            for section, label, values in rows:
                for j, v in values.items():
                    long.append({"yearbook": year, "sheet": ws.title.strip(), "output": out, "dim": dim,
                                 "title": title, "section": section, "label": label,
                                 "columns": " | ".join(paths[j]), "value": v})


# ---- Normalising ---------------------------------------------------------------------------------
INSURERS = [(r"^(a\.?ch\.?s\.?|achs|asociacion chilena de seguridad)$", "ACHS"),
            (r"^(c\.?ch\.?c\.?|museg|mutual de seguridad( c\.?ch\.?c\.?)?)$", "MUSEG"),
            (r"^(i\.?s\.?t\.?|ist|instituto de seguridad del trabajo)$", "IST"),
            (r"^(i\.?s\.?l\.?|isl|instituto de seguridad laboral)$", "ISL"),
            (r"^(sub ?- ?)?total (mutuales|mutualidades)$|^(mutuales|mutualidades) total$", "Mutual insurers")]
SEXES = {"hombres": "Male", "hombre": "Male", "mujeres": "Female", "mujer": "Female",
         "no informado": "Unknown", "sin informacion": "Unknown"}


def insurer_of(text):
    k = key(text)
    return next((v for p, v in INSURERS if re.search(p, k)), None)


def case_type(text):
    k = key(text)
    if "trayecto" in k and "trabajo" in k:
        return "Work and commuting accidents"
    if "trayecto" in k:
        return "Commuting accidents"
    if "enfermedad" in k:
        return "Occupational diseases"
    if "trabajo" in k:
        return "Work accidents"
    return None


def measure_of(title):
    k = key(title)
    if "protegid" in k:
        return ("Protected employment relationships" if "relaciones laborales" in k else
                "Protected workers (unique persons)" if "personas unicas" in k else "Protected workers")
    if k.startswith("tasa de mortalidad"):
        return "Deaths per 100,000 workers"
    if k.startswith("tasas de accidentabilidad"):
        return "Accidents per 100 workers"
    if "promedio de dias perdidos" in k:
        return "Mean days lost per case"
    if k.startswith("numero de dias perdidos"):
        return "Days lost"
    if k.startswith("numero de fallecidos"):
        return "Deaths"
    if k.startswith("numero de accidentes") or k.startswith("numero de enfermedades"):
        return "Cases"
    raise ValueError("unknown measure: " + title)


def is_total(label):
    return bool(re.match(r"^(numero )?(sub ?- ?)?total|^tasa (de )?(total|accidentes|mortalidad|accidentabilidad)",
                         key(label)))


records = []
for r in long:
    parts = [p for p in r["columns"].split(" | ") if p]
    years = [int(p) for p in parts if re.fullmatch(r"\d{4}", p.strip())]
    year = years[-1] if years else r["yearbook"]
    measure = measure_of(r["title"])
    # Insurer: from the columns, else from the row (national tables)
    insurer = next((insurer_of(p) for p in parts if insurer_of(p)), None)
    if insurer is None and any(key(p) in ("total", "total mutuales", "total mutualidades") for p in parts):
        insurer = "Mutual insurers" if any(key(p) in ("total mutuales", "total mutualidades") for p in parts) else "Total"
    sex = next((SEXES[key(p)] for p in parts if key(p) in SEXES), None)
    if sex is None and any(key(p) == "total" for p in parts) and any(key(q) in SEXES for q in parts):
        sex = "Total"
    label = r["label"]
    category = label
    if r["dim"] == "insurer":
        row_insurer = insurer_of(label) or ("Total" if is_total(label) else None)
        if key(label) in SEXES:
            sex, category = SEXES[key(label)], "Total"
        elif row_insurer:
            insurer, category = (insurer if insurer and row_insurer == "Total" and r["output"] == "sex" else row_insurer), "Total"
    elif is_total(label):
        category = "Total"
    ct = case_type(r["section"]) or case_type(label) or case_type(r["title"]) if "protegid" not in key(r["title"]) else ""
    # The 2025 yearbook leaves out COVID-19 and counts diseases by date of notification
    if ct == "Occupational diseases" and r["yearbook"] >= 2025:
        ct = "Occupational diseases (excluding COVID-19)"
    # In table 22 of the 2025 yearbook (accidents by insurer and sex), the rows of accidents for
    # MUSEG and IST are swapped: their totals belong to the other insurer in table 21.a
    if r["yearbook"] == 2025 and r["sheet"] == "22" and ct.endswith("accidents") and insurer in ("MUSEG", "IST"):
        insurer = {"MUSEG": "IST", "IST": "MUSEG"}[insurer]
    v = r["value"]
    records.append({"yearbook": r["yearbook"], "output": r["output"], "dim": r["dim"], "year": year,
                    "insurer": insurer or "Total", "sex": sex or "Total", "category_es": category,
                    "case_type": ct or "", "measure": measure, "value": v, "title": r["title"],
                    "sheet": r["sheet"]})

# Rows that are headings with numbers (they repeat totals)
records = [r for r in records if not re.match(r"^(fallecidos|accidentes) ", key(r["category_es"]))]

# A "Total" in a table of the mutual insurers only is the total of the mutual insurers. Until about
# 2020 the tables of accidents cover the mutual insurers only; later ones include the ISL
has_isl = {(r["yearbook"], r["sheet"]) for r in records if r["insurer"] == "ISL"}
for r in records:
    if r["insurer"] == "Total" and (r["yearbook"], r["sheet"]) not in has_isl and \
            re.search(r"mutual", key(r["title"])) and not re.search(r"organismo|\bisl\b", key(r["title"])):
        r["insurer"] = "Mutual insurers"

# Accident rates are fractions in some yearbooks and percentages in others
by_table = {}
for r in records:
    by_table.setdefault((r["yearbook"], r["sheet"]), []).append(r)
for rows in by_table.values():
    if rows[0]["measure"] == "Accidents per 100 workers" and max(r["value"] for r in rows) < 1:
        for r in rows:
            r["value"] *= 100

# ---- Labels --------------------------------------------------------------------------------------
labels = {}
with open(Path(__file__).resolve().parent / "labels_es_en.csv", encoding="utf-8") as fh:
    for row in csv.DictReader(fh):
        labels[(row["variable"], key(row["spanish"]))] = row
# Labels found only in SUSESO's broad groups of sectors (others, such as "Construccion", are in both)
BROAD = {"agricultura y pesca", "mineria", "comercio", "transportes y comunicaciones", "servicios",
         "servicios financieros", "servicios comunales y personales", "electricidad, gas y agua"}
missing = set()


def size_label(text):
    k = key(text)
    if k == "independientes":
        return "Self-employed"
    if k in ("sin informacion", "no informado"):
        return "Unknown"
    nums = [int(n.replace(".", "")) for n in re.findall(r"[\d.]+\d|\d", k)]
    if "y mas" in k:
        return f"{nums[0]} and over"
    return f"{nums[0]}-{nums[1]}"


for rows in by_table.values():
    cats = {key(re.sub(r"\s*\([a-z0-9]\)$", "", r["category_es"])) for r in rows}
    grouping = "Broad groups" if cats & BROAD else "Sections"
    for r in rows:
        r["category_es"] = re.sub(r"\s*\([a-z0-9]\)$", "", r["category_es"]).strip()
        r["grouping"] = grouping if r["dim"] == "sector" else ""
        if r["category_es"] == "Total" or r["dim"] == "insurer":
            r["category"] = r["category_es"]
        elif r["dim"] == "size":
            r["category"] = size_label(r["category_es"])
        else:
            lab = labels.get((r["dim"], key(r["category_es"])))
            if lab is None:
                missing.add((r["dim"], r["category_es"]))
            r["category"] = lab["english"] if lab else ""
records = [r for rows in by_table.values() for r in rows]
if missing:
    raise ValueError("No English label for: " + "; ".join(f"{d}: {t}" for d, t in sorted(missing)))

# ---- One value per year: the latest yearbook that has it ----------------------------------------
best = {}
for r in sorted(records, key=lambda r: r["yearbook"]):
    k = (r["output"], r["dim"], r["year"], r["insurer"], r["sex"], r["grouping"], key(r["category_es"]),
         r["case_type"], r["measure"])
    if k not in best or best[k]["yearbook"] < r["yearbook"]:
        best[k] = r

# A total of insurers from an older yearbook than its insurers would not add up: drop it. Tables
# that give the three mutual insurers but not their total (the tables of deaths, and the 2025
# yearbook) get the total of the mutual insurers for counts, as their sum (a blank cell is zero)
table_insurers = {}
for r in records:
    table_insurers.setdefault((r["yearbook"], r["sheet"]), set()).add(r["insurer"])
MUTUAL = ("ACHS", "IST", "MUSEG")
groups = {}
for r in best.values():
    gk = (r["output"], r["dim"], r["year"], r["sex"], r["grouping"], key(r["category_es"]),
          r["case_type"], r["measure"])
    groups.setdefault(gk, []).append(r)
final = []
for rows in groups.values():
    parts = [r for r in rows if r["insurer"] in MUTUAL + ("ISL",)]
    newest = max((r["yearbook"] for r in parts), default=0)
    kept = [r for r in rows if r in parts or r["yearbook"] >= newest]
    final += kept
    mutual = [r for r in parts if r["insurer"] in MUTUAL and r["yearbook"] == newest]
    if mutual and rows[0]["measure"] in ("Cases", "Deaths", "Days lost") and \
            set(MUTUAL) <= table_insurers[(newest, mutual[0]["sheet"])] and \
            not any(r["insurer"] == "Mutual insurers" for r in kept):
        final.append(dict(mutual[0], insurer="Mutual insurers", value=sum(r["value"] for r in mutual)))

# ---- Checks --------------------------------------------------------------------------------------
COUNTS = ("Cases", "Deaths", "Days lost")
notes = []


def check(name, rows, parts_of):
    """Share of groups whose parts add up to their total (within 0.5%)."""
    groups = {}
    for r in rows:
        if r["measure"] not in COUNTS:
            continue
        whole, part = parts_of(r)
        g = groups.setdefault(whole, [None, 0.0, 0])
        if part:
            g[1] += r["value"]
            g[2] += 1
        else:
            g[0] = r["value"]
    tested = [(t, p) for t, p, n in groups.values() if t is not None and n > 1]
    ok = sum(abs(t - p) <= 0.005 * max(t, 1) for t, p in tested)
    notes.append(f"{name}: {ok} of {len(tested)} totals equal the sum of their parts (within 0.5%)")


def base(r):
    return (r["output"], r["dim"], r["year"], r["sex"], r["case_type"], r["measure"], r["grouping"])


national = [r for r in final if r["output"] == "national" and r["sex"] == "Total"]
check("mutual insurers (national)", [r for r in national if r["insurer"] in ("ACHS", "IST", "MUSEG", "Mutual insurers")],
      lambda r: (base(r), r["insurer"] != "Mutual insurers"))
check("all insurers (national)", [r for r in national if r["insurer"] not in ("Mutual insurers",)],
      lambda r: (base(r), r["insurer"] != "Total"))
# The tables by sex give each insurer's total: it should equal the national table
# (in the same yearbook, since later yearbooks revise)
nat = {(r["yearbook"], r["year"], r["insurer"], r["case_type"], r["measure"]): r["value"] for r in records
       if r["output"] == "national" and r["sex"] == "Total"}
pairs = [(r["value"], nat[k], k) for r in records if r["output"] == "sex" and r["dim"] == "insurer" and r["sex"] == "Total"
         for k in [(r["yearbook"], r["year"], r["insurer"], r["case_type"], r["measure"])]
         if k in nat and r["measure"] in COUNTS]
off = [k for a, b, k in pairs if abs(a - b) > 0.005 * max(b, 1)]
notes.append(f"insurers by sex: {len(pairs) - len(off)} of {len(pairs)} totals equal the national tables "
             f"(within 0.5%); not equal: {off}")
for out in ("sector", "region", "size"):
    check(f"{out}s", [r for r in final if r["output"] == out],
          lambda r: (base(r) + (r["insurer"],), r["category"] != "Total"))
check("sexes", [r for r in final if r["output"] == "sex"],
      lambda r: ((r["dim"], r["year"], r["insurer"], key(r["category_es"]), r["case_type"], r["measure"]),
                 r["sex"] != "Total"))


# ---- Write ----------------------------------------------------------------------------------------
def write(name, rows, columns):
    rows = sorted(rows, key=lambda r: tuple(str(r.get(c, "")) for c in columns))
    with open(tidy / name, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


tail = ["case_type", "measure", "value", "source_yearbook"]
for r in final:
    r["source_yearbook"] = r["yearbook"]
    r["by"] = r["dim"].capitalize()
    name = {"sector": "sector", "region": "region", "size": "workers"}.get(r["dim"])
    if name and r["output"] != "sex":
        r[name], r[name + "_es"] = r["category"], r["category_es"]
write("national.csv", [r for r in final if r["output"] == "national"], ["year", "insurer", "sex"] + tail)
write("by_sector.csv", [r for r in final if r["output"] == "sector"],
      ["year", "grouping", "sector", "sector_es", "insurer"] + tail)
write("by_region.csv", [r for r in final if r["output"] == "region"],
      ["year", "region", "region_es", "insurer"] + tail)
write("by_size.csv", [r for r in final if r["output"] == "size"],
      ["year", "workers", "workers_es", "insurer"] + tail)
write("by_sex.csv", [r for r in final if r["output"] == "sex"],
      ["year", "sex", "by", "category", "category_es", "insurer"] + tail)
for f in sorted(tidy.glob("*.csv")):
    print(f"{f.name:16} {sum(1 for _ in open(f, encoding='utf-8')) - 1:7} rows")
print("\n".join(notes))
