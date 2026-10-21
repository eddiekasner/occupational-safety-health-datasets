"""Extract tables from the Labour Department's "Occupational Safety and Health Statistics" (PDF).

The Labour Department publishes Hong Kong's accident and occupational disease statistics each year
as a bilingual PDF. This script reads two editions (downloaded by scripts/fetch_data.R):
  - 2025: ten-year tables (2016 to 2025) and tables for 2024 and 2025 by industry and type of
    accident
  - 2015: ten-year tables (2006 to 2015, and industrial accidents from 1999)
and writes, to data/tidy/:
  trends.csv                  national and major-industry series by year: industrial accidents,
                              occupational injuries, fatalities, workers and rates
  occupational_diseases.csv   confirmed occupational disease cases by disease, 2006 to 2025
  by_industry_2024_2025.csv   industrial accidents and occupational injuries by industry
  by_type_2024_2025.csv       the same by industry and type of accident, with fatalities
  health_services.csv         the Labour Department's occupational health services by year
  clinic_patients.csv         new patients at the Labour Department's occupational health clinics
                              by disease group

Tables are read from the position of each word on the page: numbers are assigned to the nearest
year column (Hong Kong writes thousands with a space, "3 720"), Chinese labels are dropped, and
fatalities printed in brackets, "(5)", are kept separately.

The two Excel workbooks, from data.gov.hk, give the Labour Department's occupational health
services and clinic patients.

Needs pdfplumber and openpyxl (pip install pdfplumber openpyxl) and pdftotext (poppler). Written in Python because no R package reads these
tables reliably; the outputs are committed, so you only need to run this to update them.

Usage (from the module folder):  python scripts/extract_tables.py
"""
import csv
import re
import subprocess
from pathlib import Path

import openpyxl
import pdfplumber

module = Path(__file__).resolve().parent.parent
raw = module / "data" / "raw"
tidy = module / "data" / "tidy"

SECTION = re.compile(r"^(No\. of|Number of|Accident rate|Injury rate|Fatality rate)", re.I)
STOP = re.compile(r"^(Notes?:|\u8a3b\u91cb)")


def page_lines(page):
    """Words of a page grouped into lines (by vertical position), each sorted left to right."""
    rows = {}
    for w in page.extract_words(x_tolerance=1.5):
        rows.setdefault(round(w["top"] / 3), []).append(w)
    return [sorted(v, key=lambda w: w["x0"]) for _, v in sorted(rows.items())]


def ascii_word(w):
    return all(ord(c) < 128 for c in w["text"])


def parse_page(page):
    """Return (title, columns, rows); rows are dicts with section, label and {column: (value,
    fatalities)}."""
    lines = page_lines(page)
    title = " ".join(w["text"] for w in lines[0] if ascii_word(w))
    # The header is the line with the most years (titles also contain years, "2006 - 2015")
    header, cols = None, []
    for i, ln in enumerate(lines):
        years = [w for w in ln if re.fullmatch(r"(19|20)\d\d", w["text"])]
        if len(years) >= 2 and len(years) > len(cols):
            header = i
            cols = [(w["text"], (w["x0"] + w["x1"]) / 2) for w in years]
    if header is None:
        return title, [], []
    first_x = min(x for _, x in cols)
    rows, pending, section, prev_open = [], "", "", False
    for ln in lines[header + 1:]:
        words = [w for w in ln if ascii_word(w)]
        if not words:
            continue
        text_words = [w for w in words if w["x1"] < first_x - 25 and not re.fullmatch(r"\d+", w["text"])]
        nums = [w for w in words if w["x0"] >= first_x - 40 and
                re.fullmatch(r"\(?\d+(\.\d+)?\)?", w["text"])]
        label = " ".join(w["text"] for w in text_words).strip()
        if STOP.match(label):
            break
        if not nums:
            if not label:
                continue
            if SECTION.match(label):
                section, pending = label, ""
            elif rows and (prev_open or label[:1].islower() or label.startswith("(")):
                rows[-1]["label"] += " " + label
            else:
                pending = (pending + " " + label).strip()
            prev_open = False
            continue
        values = {}
        for w in nums:
            x = (w["x0"] + w["x1"]) / 2
            col = min(cols, key=lambda c: abs(c[1] - x))[0]
            v, f = values.get(col, ("", ""))
            if w["text"].startswith("("):
                f = w["text"].strip("()")
            else:
                v += w["text"]
            values[col] = (v, f)
        full = (pending + " " + label).strip()
        prev_open = not label and bool(pending)
        rows.append({"section": section, "label": full, "values": values})
        pending = ""
    return title, [c for c, _ in cols], rows


def number(s):
    if s in ("", None) or s.upper() == "N.A.":
        return ""
    return float(s) if "." in s else int(s)


def subtitle(page):
    """The "analysed by ..." line under a page's title, to tell apart pages with the same title."""
    for ln in page_lines(page)[:6]:
        text = " ".join(w["text"] for w in ln if ascii_word(w))
        if "analysed by" in text:
            return text.strip("- ").replace("analysed by ", "")
    return ""


# ---- Trends: national and major-industry series ----------------------------------------------
# Measures, recognised from the row and section labels of the ten-year tables
measures = [
    (r"^No\. of Accidents", "Industrial accidents"),
    (r"^No\. of Fatalities", "Industrial accident fatalities"),
    (r"^(Employment Size|Total No\. of Workers)", "Workers (employment size)"),
    (r"^Acc(\.|ident) rate", "Industrial accident rate per 1,000 workers"),
    (r"^Fatality rate per workers", "Industrial accident fatality rate per 1,000 workers"),
    (r"^Fatality rate/", "Industrial accident fatality rate per 1,000 workers"),
    (r"^Number of occupational injur", "Occupational injuries"),
    (r"^Injury rate", "Occupational injury rate per 1,000 employees"),
    (r"^Number of work fatalities", "Work fatalities"),
    (r"^Fatality rate per employees", "Work fatality rate per 1,000 employees"),
]
industries = [("Construction", "Construction"), ("Catering", "Catering (to 2008)"),
              ("Food and Beverage", "Food and beverage services (from 2009)"),
              ("Manufacturing", "Manufacturing"), ("All Industries", "All industries"),
              ("All Economic", "All economic activities")]

trends, diseases = [], []
for edition, pages in [(2015, [36, 37, 38, 39, 40, 43, 44, 46]), (2025, [36, 37, 38, 40, 41, 42])]:
    # (The "Major Industries" pages repeat the industry pages, so they are skipped)
    pdf = pdfplumber.open(raw / f"OSH_Statistics_{edition}_en.pdf")
    for n in pages:
        title, cols, rows = parse_page(pdf.pages[n - 1])
        industry = next(v for k, v in industries if k in title)
        for r in rows:
            name = (r["label"] or r["section"]).strip()
            measure = next((m for pat, m in measures if re.match(pat, name, re.I)), None)
            if measure is None:
                raise ValueError(f"{edition} page {n}: unknown row {name!r}")
            for year, (v, f) in r["values"].items():
                trends.append({"year": int(year), "measure": measure, "industry": industry,
                               "value": number(v), "source_edition": edition})

# ---- Occupational diseases --------------------------------------------------------------------
disease_names = {
    "Compressed Air Illness": "Compressed air illness",
    "Tenosynovitis of the Hand or Forearm": "Tenosynovitis of the hand or forearm",
    "Carpal Tunnel Syndrome": "Carpal tunnel syndrome", "Beat Knee": "Beat knee",
    "Tuberculosis": "Tuberculosis", "Streptococcus suis Infection": "Streptococcus suis infection",
    "Cadmium Poisoning": "Cadmium poisoning", "Occupational Dermatitis": "Occupational dermatitis",
    "Poisoning by Halogen": "Poisoning by halogen derivatives of aliphatic hydrocarbons",
    "Inflammation or Ulceration": "Inflammation or ulceration of the upper respiratory tract or "
                                  "mouth (dust, liquid or vapour)",
    "Brucellosis": "Brucellosis", "Phosphorous Poisoning": "Phosphorus poisoning",
    "Leptospirosis": "Leptospirosis", "Avian chlamydiosis": "Avian chlamydiosis",
    "Occupational Asthma": "Occupational asthma", "Silicosis": "Silicosis",
    "Asbestosis": "Asbestosis", "Mesothelioma": "Mesothelioma",
    "Occupational Deafness": "Occupational deafness", "Legionnaires": "Legionnaires' disease",
    "Gas Poisoning": "Gas poisoning", "Mercury Poisoning": "Mercury poisoning", "Total": "Total",
}


def disease_name(label):
    for k, v in disease_names.items():
        if label.startswith(k):
            return v
    raise ValueError("unknown disease: " + label)


# 2016 to 2025, from the word positions
pdf = pdfplumber.open(raw / "OSH_Statistics_2025_en.pdf")
title, cols, rows = parse_page(pdf.pages[42])
for r in rows:
    for year, (v, f) in r["values"].items():
        diseases.append({"year": int(year), "disease": disease_name(r["label"]),
                         "cases": number(v), "incidents": number(f), "source_edition": 2025})

# 2006 to 2015: this page's labels wrap over several lines, so read it from the pdftotext layout,
# where each row's numbers sit on one line, in the order of the table
order_2015 = ["Compressed Air Illness", "Tenosynovitis of the Hand or Forearm",
              "Carpal Tunnel Syndrome", "Tuberculosis", "Streptococcus suis Infection",
              "Cadmium Poisoning", "Occupational Dermatitis", "Poisoning by Halogen",
              "Inflammation or Ulceration", "Brucellosis", "Phosphorous Poisoning", "Leptospirosis",
              "Avian chlamydiosis", "Occupational Asthma", "Silicosis", "Asbestosis",
              "Mesothelioma", "Occupational Deafness", "Gas Poisoning", "Total"]
text = subprocess.run(["pdftotext", "-layout", "-f", "45", "-l", "45",
                       str(raw / "OSH_Statistics_2015_en.pdf"), "-"],
                      capture_output=True, check=True).stdout.decode("utf-8")
value = r"(N\.A\.|\d+(?: \(\d+\))?)"
found = [re.findall(value, m.group(1)) for m in
         re.finditer(r"((?:\s+" + value + r"){10})\s*$", text, re.M)]
years_2015 = list(range(2006, 2016))
found = [f for f in found if len(f) == 10 and f != [str(y) for y in years_2015]]  # not the header
if len(found) != len(order_2015):
    raise ValueError(f"2015 diseases: expected {len(order_2015)} rows, found {len(found)}")
for name, vals in zip(order_2015, found):
    for year, v in zip(years_2015, vals):
        m = re.match(r"(\d+) \((\d+)\)", v)
        diseases.append({"year": year, "disease": disease_name(name),
                         "cases": number(m.group(1) if m else v),
                         "incidents": number(m.group(2)) if m else "", "source_edition": 2015})

# Check: the diseases add up to the total in every year
for year in sorted({d["year"] for d in diseases}):
    parts = sum(d["cases"] or 0 for d in diseases if d["year"] == year and d["disease"] != "Total")
    total = [d["cases"] for d in diseases if d["year"] == year and d["disease"] == "Total"][0]
    if parts != total:
        raise ValueError(f"diseases {year}: parts {parts} != total {total}")

# ---- 2024 and 2025 by industry and type of accident --------------------------------------------
# Manufacturing industries reported separately; they add up to Manufacturing
manufacturing = {"Textiles", "Wearing apparel", "Electronics", "Fabricated metal products",
                 "Rubber and plastics products", "Shipbuilding and shiprepairing",
                 "Printing and reproduction of recorded media", "Other manufacturing industries"}
by_type, by_industry = [], []
pdf = pdfplumber.open(raw / "OSH_Statistics_2025_en.pdf")
for n in range(4, 34):
    page = pdf.pages[n - 1]
    title, cols, rows = parse_page(page)
    statistic = "Industrial accidents" if title.startswith("Industrial") else "Occupational injuries"
    industry = re.sub(r"^(Industrial Accidents|Occupational Injuries) in ", "", title)
    industry = re.sub(r" (Industry|Undertakings)$", "", industry)
    industry = "Total" if industry.startswith("All") else industry[0] + industry[1:].lower()
    if industry == "Professional and business services":
        # This page leaves out pest control and cleaning services, which have their own page
        industry += " (excluding pest control and cleaning)"
    analysed = subtitle(page)
    if analysed in ("Industry", "Industry Section"):
        labels = [r["label"] for r in rows]
        if n == 4:
            # Two labels wrap in this table; its rows are always in this order
            labels = ["Manufacturing", "Textiles", "Wearing apparel", "Electronics",
                      "Fabricated metal products", "Rubber and plastics products",
                      "Shipbuilding and shiprepairing",
                      "Printing and reproduction of recorded media",
                      "Other manufacturing industries", "Construction",
                      "Food and beverage services", "Other industries", "Total"]
            if len(labels) != len(rows):
                raise ValueError("page 4 has a different number of rows")
        for label, r in zip(labels, rows):
            label = re.sub(r"^[IV]+\. ", "", label).replace("TOTAL", "Total")
            for year, (v, f) in r["values"].items():
                by_industry.append({"statistic": statistic, "industry": label,
                                    "parent": "Manufacturing" if label in manufacturing else "",
                                    "year": int(year), "cases": number(v),
                                    "fatalities": number(f) if f else 0})
    else:
        for r in rows:
            for year, (v, f) in r["values"].items():
                parent = "Manufacturing" if industry in manufacturing else ""
                by_type.append({"statistic": statistic, "industry": industry, "parent": parent,
                                "accident_type": r["label"].replace("TOTAL", "Total"),
                                "year": int(year), "cases": number(v),
                                "fatalities": number(f) if f else 0})

# Check: the manufacturing industries add up to Manufacturing
for year in (2024, 2025):
    rows = [r for r in by_industry if r["statistic"] == "Industrial accidents" and r["year"] == year]
    if sum(r["cases"] for r in rows if r["parent"]) != \
            [r["cases"] for r in rows if r["industry"] == "Manufacturing"][0]:
        raise ValueError(f"manufacturing industries do not add up in {year}")

# Check: each industry's types add up to its total, and the by-industry tables match the totals of
# the industry pages with the same name
for key in {(r["statistic"], r["industry"], r["year"]) for r in by_type}:
    rows = [r for r in by_type if (r["statistic"], r["industry"], r["year"]) == key]
    total = [r["cases"] for r in rows if r["accident_type"] == "Total"][0]
    if sum(r["cases"] for r in rows if r["accident_type"] != "Total") != total:
        raise ValueError(f"types do not add up for {key}")
    match = [r["cases"] for r in by_industry
             if (r["statistic"], r["industry"], r["year"]) == key]
    if key[1].startswith("Professional and business"):
        match = [r["cases"] for r in by_industry if r["statistic"] == key[0] and r["year"] == key[2]
                 and r["industry"] == "Professional and business services"]
        total = sum(r["cases"] for r in by_type if r["statistic"] == key[0] and r["year"] == key[2]
                    and r["accident_type"] == "Total" and
                    (r["industry"].startswith("Professional") or r["industry"].startswith("Pest")))
    if match and match[0] != total:
        raise ValueError(f"industry total differs for {key}: {match[0]} vs {total}")

# ---- Occupational health services (data.gov.hk workbooks) --------------------------------------
services, patients = [], []
wb = openpyxl.load_workbook(raw / "ohs-keystats_en.xlsx", data_only=True)
rows = [r for r in wb.worksheets[0].iter_rows(values_only=True) if r[0] is not None]
for r in rows[1:]:
    for activity, v in zip(rows[0][1:], r[1:]):
        services.append({"year": int(r[0]), "activity": activity, "count": int(v)})
wb = openpyxl.load_workbook(
    raw / "ohs-keystats-new-patients-attend-ohc-breakdown-type-of-diseases_en.xlsx", data_only=True)
rows = [r for r in wb.worksheets[0].iter_rows(values_only=True) if r[0] is not None]
for r in rows[1:]:
    for group, v in zip(rows[0][1:], r[1:]):
        if group:
            group = group.replace("Number of ", "").replace("Others", "Other")
            patients.append({"year": int(r[0]), "disease_group": group.replace("Diseases", "diseases"),
                             "new_patients": int(v)})


def write(name, rows):
    with open(tidy / name, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


write("trends.csv", sorted(trends, key=lambda r: (r["measure"], r["industry"], r["year"])))
write("occupational_diseases.csv", sorted(diseases, key=lambda r: (r["year"], r["disease"])))
write("by_type_2024_2025.csv", by_type)
write("by_industry_2024_2025.csv", by_industry)
write("health_services.csv", services)
write("clinic_patients.csv", patients)
for f in sorted(tidy.glob("*.csv")):
    print(f"{f.name:28} {sum(1 for _ in open(f)) - 1:6} rows")
